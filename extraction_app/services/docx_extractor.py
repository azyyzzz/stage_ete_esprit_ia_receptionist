"""
Extraction Word (.docx) -- adapte de data/scripts/extract_admission_docx.py.

Regroupe le contenu par SECTION (dernier titre en gras vu dans le document),
pas par question individuelle : une fiche qui regroupe toutes les infos d'un
meme sujet donne une reponse plus complete a une question vague qu'une fiche
par mini-question isolee, et remonte mieux en recherche semantique.

A l'interieur d'une section, detecte les paires question/reponse (question =
paragraphe finissant par "?"), et accumule le reste comme texte descriptif.
"""

from __future__ import annotations

import re
from pathlib import Path

import docx

PLACEHOLDER_PATTERN = re.compile(r"\((?:mettre|bouton|cliquable)[^)]*\)", re.IGNORECASE)


def _slugify(text: str) -> str:
    text = re.sub(r"[^a-zA-Z0-9]+", "_", text.strip().lower())
    return text.strip("_")[:60]


def _is_bold_paragraph(paragraph) -> bool:
    if not paragraph.runs:
        return False
    return all(run.bold for run in paragraph.runs if run.text.strip())


def _is_question(text: str) -> bool:
    return text.strip().endswith("?")


def _strip_numbering(text: str) -> str:
    return re.sub(r"^\d+[.)]\s*", "", text).strip()


def _strip_author_placeholders(text: str) -> str:
    return re.sub(r"\s+", " ", PLACEHOLDER_PATTERN.sub("", text)).strip()


def _group_into_sections(paragraphs: list) -> list[dict]:
    sections: list[dict] = []
    current_title = None
    current_lines: list[str] = []

    def flush():
        if current_title is not None and current_lines:
            sections.append({"titre": current_title, "lignes": list(current_lines)})

    for paragraph in paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        if _is_bold_paragraph(paragraph):
            flush()
            current_title = text
            current_lines = []
            continue

        if current_title is None:
            continue  # texte avant le premier titre en gras (ex. titre du document)

        if "\n" in text:
            first_line, rest = text.split("\n", 1)
            first_line = _strip_numbering(first_line)
            if _is_question(first_line):
                rest = _strip_author_placeholders(rest.strip())
                current_lines.append(f"Q: {first_line}\nR: {rest}")
                continue

        text = _strip_numbering(text)
        text = _strip_author_placeholders(text)
        if not text:
            continue

        if _is_question(text):
            current_lines.append(f"Q: {text}")
        elif current_lines and current_lines[-1].startswith("Q:") and "\nR:" not in current_lines[-1]:
            current_lines[-1] += f"\nR: {text}"
        elif current_lines and current_lines[-1].split("\n")[-1].startswith("R:"):
            current_lines[-1] += f" {text}"
        elif current_lines and current_lines[-1].startswith("- "):
            current_lines[-1] += f" {text}"
        else:
            current_lines.append(f"- {text}")

    flush()
    return sections


def extract_docx(docx_path: Path, categorie: str, source: str, id_slug: str | None = None) -> list[dict]:
    """Retourne une liste de fiches {id, categorie, titre, contenu, source}.
    Voir pdf_extractor.extract_pdf pour le role de id_slug."""
    document = docx.Document(str(docx_path))
    sections = _group_into_sections(document.paragraphs)

    slug = _slugify(id_slug if id_slug is not None else docx_path.stem)
    doc_label = id_slug if id_slug is not None else docx_path.stem

    records = []
    for i, section in enumerate(sections, start=1):
        content = "\n\n".join(section["lignes"]).strip()
        if len(content) < 30:
            continue
        records.append(
            {
                "id": f"upload_docx_{slug}_{i:02d}",
                "categorie": categorie,
                "titre": f"{doc_label} — {section['titre']}",
                "contenu": content,
                "source": source,
            }
        )
    return records
